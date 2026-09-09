"""Build the reconciliation workbook from the API's own report document.

One responsibility: turn the JSON document that /api/v1/report/{date} returns
into a formatted .xlsx, column-per-field. Taking the same document as input is
deliberate - the spreadsheet and the JSON cannot disagree about a figure.

Three sheets:
  Summary  the headline block, one figure per row
  Sources  deduplicated, one row per vendor      - what you report
  Runs     every raw table row, all columns      - what you prove it from
"""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

TEAL = "137D73"
GREEN = "22885A"
RED = "A8483D"
AMBER = "94590A"
GREY = "7E8F9A"
INK = "1B2B38"
MONEY = '#,##0.00'
COUNT = "#,##0"
PCT = '0.00"%"'
THIN = Side(style="thin", color="E4EAEC")

#  (header, dotted path into the JSON, number format)
SOURCE_COLS = [
    ("Vendor ID", "vendorId", "@"), ("Bank / source", "name", None),
    ("Status", "status", None),
    ("Records (partner)", "records.partner", COUNT),
    ("Records (SAP)", "records.sap", COUNT),
    ("Matched", "records.matched", COUNT),
    ("Unmatched", "records.unmatched", COUNT),
    ("Amount partner", "amount.partner", MONEY),
    ("Amount SAP", "amount.sap", MONEY),
    ("Balance difference", "amount.difference", MONEY),
    ("Anomalies", "anomalies.total", COUNT),
    ("Missing in SAP", "anomalies.byType.missingInSap", COUNT),
    ("Missing in partner", "anomalies.byType.missingInPartner", COUNT),
    ("SAP amount high", "anomalies.byType.sapAmountHigh", COUNT),
    ("Partner amount high", "anomalies.byType.partnerAmountHigh", COUNT),
    ("Duplicates", "anomalies.byType.duplicates", COUNT),
    ("Date differences", "anomalies.byType.dateDifferences", COUNT),
    ("Match rate %", "rates.matchPct", PCT),
    ("Anomaly rate %", "rates.anomalyPct", PCT),
    ("Value unconfirmed %", "rates.valueUnconfirmedPct", PCT),
    ("Runs", "runs.count", COUNT),
    ("Runs superseded", "runs.superseded", COUNT),
    ("Figures varied", "runs.figuresVariedAcrossRuns", None),
    ("Row key", "provenance.rowKey", "@"),
    ("Partition key", "provenance.partitionKey", "@"),
    ("Upload date", "provenance.uploadDate", "@"),
    ("Generated date", "provenance.generatedDate", "@"),
    ("Exposure", "exposure", MONEY),
    ("Pipeline assessment", "assessment", None),
]

RUN_COLS = [
    ("Authoritative", "isAuthoritative", None),
    ("Superseded by", "supersededBy", "@"),
    ("Vendor ID", "vendorId", "@"), ("Bank / source", "sourceName", None),
    ("Business date", "businessDate", "@"),
    ("Upload date/time", "uploadDateTime", "@"),
    ("Generated date/time", "generatedDateTime", "@"),
    ("Service timestamp", "timestamp", "@"),
    ("Records (partner)", "totalRecordsPartner", COUNT),
    ("Records (SAP)", "totalRecordsSap", COUNT),
    ("Matched", "m00Matched", COUNT),
    ("Amount partner", "totalAmountPartner", MONEY),
    ("Amount SAP", "totalAmountSap", MONEY),
    ("Balance difference", "balanceDifference", MONEY),
    ("Anomalies", "totalAnomalies", COUNT),
    ("Missing in SAP", "missingInSap", COUNT),
    ("Missing in partner", "missingInPartner", COUNT),
    ("SAP amount high", "sapAmountHigh", COUNT),
    ("Partner amount high", "partnerAmountHigh", COUNT),
    ("Duplicates", "duplicates", COUNT),
    ("Date differences", "dateDifferences", COUNT),
    ("Match rate %", "matchRatePct", PCT),
    ("Anomaly rate %", "anomalyRatePct", PCT),
    ("Status", "status", None),
    ("Row key", "rowKey", "@"), ("Partition key", "partitionKey", "@"),
    ("Upload date (raw)", "uploadDate", "@"),
    ("Generated date (raw)", "generatedDate", "@"),
    ("Exposure", "exposure", MONEY),
    ("Pipeline assessment", "generalSummary", None),
]


def _dig(obj: dict, path: str):
    for part in path.split("."):
        if obj is None:
            return None
        obj = obj.get(part)
    return obj


def _sheet(wb, title, cols, rows, accent=TEAL):
    ws = wb.create_sheet(title)
    ws.append([c[0] for c in cols])
    for r in rows:
        ws.append([_dig(r, path) for _, path, _ in cols])

    fill = PatternFill("solid", fgColor=accent)
    for i in range(1, len(cols) + 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34

    for r in range(2, ws.max_row + 1):
        for i, (header, _, fmt) in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=i)
            if fmt:
                cell.number_format = fmt
            if fmt == "@":
                cell.alignment = Alignment(horizontal="left")
            cell.border = Border(bottom=THIN)
            v = cell.value
            if header == "Status":
                cell.font = Font(bold=True, color=GREEN if v == "Reconciled" else RED)
                cell.alignment = Alignment(horizontal="center")
            elif header in ("Anomalies", "Balance difference") and v:
                cell.font = Font(bold=True, color=RED)
            elif header == "Authoritative":
                cell.value = "Yes" if v else "No"
                cell.font = Font(bold=True, color=GREEN if v else GREY)
                cell.alignment = Alignment(horizontal="center")
            elif header == "Figures varied":
                cell.value = "YES - VERIFY" if v else "no"
                if v:
                    cell.font = Font(bold=True, color=RED)
            elif header == "Runs" and (v or 0) > 1:
                cell.font = Font(bold=True, color=AMBER)
            elif header == "Pipeline assessment":
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    tbl = Table(displayName=f"{title}Data", ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tbl)

    for i, (header, _, fmt) in enumerate(cols, start=1):
        width = 90 if header == "Pipeline assessment" else \
                34 if "key" in header.lower() else \
                26 if header in ("Bank / source",) else \
                max(len(header) + 3, 12)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "C2"
    return ws


def _summary_sheet(wb, doc):
    """Headline figures, one per row - the block a reader sees first."""
    ws = wb.create_sheet("Summary", 0)
    s, d = doc["summary"], doc["dedup"]
    day = date.fromisoformat(doc["date"])
    bad = s["sources"]["notReconciled"]

    ws["A1"] = "Bank Reconciliation"
    ws["A1"].font = Font(bold=True, size=17, color=INK)
    ws["A2"] = day.strftime("%d %B %Y")
    ws["A2"].font = Font(size=11, color=GREY)
    ws["A3"] = (f"{bad} of {s['sources']['total']} sources did not reconcile"
                if bad else f"All {s['sources']['total']} sources reconciled")
    ws["A3"].font = Font(bold=True, size=13, color=RED if bad else GREEN)

    blocks = [
        ("SOURCES", [
            ("Total sources", s["sources"]["total"], COUNT),
            ("Reconciled", s["sources"]["reconciled"], COUNT),
            ("Not reconciled", s["sources"]["notReconciled"], COUNT)]),
        ("RECORDS", [
            ("Partner file", s["records"]["partner"], COUNT),
            ("SAP", s["records"]["sap"], COUNT),
            ("Matched", s["records"]["matched"], COUNT),
            ("Unmatched", s["records"]["unmatched"], COUNT)]),
        ("AMOUNT (AED)", [
            ("Partner file", s["amount"]["partner"], MONEY),
            ("SAP", s["amount"]["sap"], MONEY),
            ("Balance difference", s["amount"]["difference"], MONEY),
            ("Reconciled value", s["amount"]["reconciled"], MONEY)]),
        ("ANOMALIES", [("Total", s["anomalies"]["total"], COUNT)]
            + [(k[0].upper() + "".join(" " + c.lower() if c.isupper() else c for c in k[1:]),
                v, COUNT) for k, v in s["anomalies"]["byType"].items()]),
        ("RATES", [
            ("Transaction match", s["rates"]["transactionMatchPct"], PCT),
            ("Source completion", s["rates"]["sourceCompletionPct"], PCT),
            ("Value reconciled", s["rates"]["valueReconciledPct"], PCT)]),
        ("DEDUPLICATION", [
            ("Strategy", d["strategy"], None),
            ("Runs collapsed", d["runsCollapsed"], COUNT),
            ("Raw runs in file", doc.get("runCount", 0), COUNT)]),
    ]

    row = 5
    for heading, items in blocks:
        ws.cell(row=row, column=1, value=heading).font = Font(bold=True, size=9, color=TEAL)
        row += 1
        for label, value, fmt in items:
            ws.cell(row=row, column=1, value=label).font = Font(color=GREY, size=10)
            c = ws.cell(row=row, column=2, value=value)
            c.font = Font(bold=True, size=11)
            if fmt:
                c.number_format = fmt
            row += 1
        row += 1

    note = ws.cell(row=row, column=1)
    varied = d.get("sourcesWithVariedFigures") or []
    if varied:
        note.value = ("WARNING: " + ", ".join(varied) + " produced different figures "
                      "across re-runs. The latest run is shown - verify which is "
                      "authoritative before reporting these numbers.")
        note.font = Font(bold=True, color=RED, size=10)
    elif d["runsCollapsed"]:
        note.value = (f"{d['runsCollapsed']} duplicate run(s) collapsed. The source table "
                      "logs every reconciliation run; the Runs sheet holds them all, with "
                      "the superseded ones marked.")
        note.font = Font(italic=True, color=AMBER, size=10)
    else:
        note.value = "No re-runs on this date. Each source reported once."
        note.font = Font(italic=True, color=GREY, size=10)

    ws.cell(row=row + 2, column=1, value=f"Generated {doc['generatedAt']}").font = \
        Font(color=GREY, size=9)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 26
    return ws


def build(doc: dict) -> BytesIO:
    """Workbook for one business date, from the report document."""
    wb = Workbook()
    wb.remove(wb.active)
    _summary_sheet(wb, doc)
    _sheet(wb, "Sources", SOURCE_COLS, doc.get("sources", []))
    if doc.get("runs"):
        _sheet(wb, "Runs", RUN_COLS, doc["runs"], accent=INK)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def filename(day) -> str:
    d = day if isinstance(day, str) else day.isoformat()
    return f"reconciliation_{d}.xlsx"
